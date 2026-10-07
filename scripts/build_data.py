#!/usr/bin/env python3
"""Audita los motores y genera un contrato compacto para el dashboard."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

MONTHS = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12,
}
MONTH_LABELS = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
MAX_XLSX_UNCOMPRESSED = 180 * 1024 * 1024


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    ascii_text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^A-Z0-9]+", " ", ascii_text.upper()).strip()


def clean_cc(value: Any) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"): text = text[:-2]
    digits = re.sub(r"\D", "", text)
    return digits.zfill(5) if 1 <= len(digits) <= 5 else ""


def clean_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_date(value: Any, epoch=None) -> date | None:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    text = str(value if value is not None else "").strip()
    if not text or normalize(text) in {"NA", "N A", "NULL", "NONE"}: return None
    text = re.sub(r"\s*([/-])\s*", r"\1", text)
    try:
        if re.fullmatch(r"\d{4,5}(?:\.\d+)?", text):
            result = from_excel(float(text), **({"epoch": epoch} if epoch else {}))
            return result.date() if isinstance(result, datetime) else None
        for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
            try: return datetime.strptime(text, pattern).date()
            except ValueError: pass
        return datetime.fromisoformat(text).date()
    except (ValueError, OverflowError): return None


def as_iso(value: Any) -> str | None:
    parsed = parse_date(value)
    return parsed.isoformat() if parsed else None


def resolve_columns(headers: Iterable[Any], rules: dict[str, tuple[str, ...]], required: set[str], label: str) -> dict[str, int]:
    headers = [clean_header(value) for value in headers]
    indexed = defaultdict(list)
    for index, header in enumerate(headers):
        if header: indexed[normalize(header)].append(index)
    result = {}
    for name, aliases in rules.items():
        matches = {index for alias in aliases for index in indexed[normalize(alias)]}
        if len(matches) > 1: raise ValueError(f"{label}: columnas ambiguas para {name}")
        if matches: result[name] = next(iter(matches))
        elif name in required: raise ValueError(f"{label}: falta columna {name} ({' / '.join(aliases)})")
    return result


DIRECTORY_COLUMNS = {
    "cc": ("CC",), "store": ("Tienda", "CC Nombre"), "status": ("Estatus", "Status"),
    "region": ("Región",), "dm": ("DM",), "rd": ("Director Regional",), "division": ("División",),
    "opened": ("Fecha de Apertura",), "city": ("Ciudad",), "state": ("Estado",), "address": ("Dirección",),
    "format": ("Formato de Tienda (Tipo tienda 1)",), "generator": ("Generador (Tipo tienda 2)",),
    "family": ("Familia (Tipo tienda 3)",), "type5": ("Tipo tienda 5",), "design": ("Intensión del Diseño",),
    "size": ("Tamaño",), "tier": ("TIER",), "seats": ("# Seats in Store",), "manager": ("Gerente",),
    "storeEmail": ("Mail de Tienda",), "managerEmail": ("Mail Gerente",),
    "appName": ("Nombre APP y Signage",), "districtCode": ("No. de Región",),
}


def load_directory(path: Path) -> tuple[list[dict], dict[str, set[str]], dict]:
    rows = load_csv(path)
    if not rows: raise ValueError("Directorio vacío")
    headers = list(rows[0])
    columns = resolve_columns(headers, DIRECTORY_COLUMNS, {"cc", "store", "status"}, "Directorio")
    directory = []; aliases = defaultdict(set); seen = set(); excluded = Counter(); excluded_cc = set(); invalid_dates = 0
    for line, row in enumerate(rows, start=2):
        values = list(row.values()); mapped = {name: values[index] for name, index in columns.items()}
        cc = clean_cc(mapped["cc"]); status = normalize(mapped["status"])
        for alias in (mapped.get("store"), mapped.get("appName")):
            if cc and normalize(alias): aliases[normalize(alias)].add(cc)
        if status != "ABIERTA":
            excluded[clean_header(mapped["status"]) or "Sin Estatus"] += 1
            if cc: excluded_cc.add(cc)
            continue
        if not cc: raise ValueError(f"Directorio fila {line}: CC inválido en tienda abierta")
        if cc in seen: raise ValueError(f"Directorio: CC abierto duplicado {cc}")
        seen.add(cc)
        item = {name: clean_header(mapped.get(name)) for name in DIRECTORY_COLUMNS if name not in {"appName", "districtCode"}}
        item.update(cc=cc, status="Abierta", opened=as_iso(mapped.get("opened")), seats=number(mapped.get("seats")))
        if not item["store"]: raise ValueError(f"Directorio fila {line}: nombre vacío en tienda abierta")
        if clean_header(mapped.get("opened")) and not item["opened"]: invalid_dates += 1
        directory.append(item)
    if not directory: raise ValueError("Directorio sin tiendas con Estatus Abierta")
    return directory, aliases, {"rows":len(rows), "validStores":len(directory), "statusFilter":"Abierta",
        "excludedStatuses":dict(excluded), "excludedStores":sum(excluded.values()), "excludedCecos":sorted(excluded_cc-seen),
        "invalidOpeningDates":invalid_dates, "headers":{name:headers[index] for name,index in columns.items()}}


QUERY_COLUMNS = {
    "employee": ("NUM_EMP", "Número de empleado", "Numero empleado"),
    "cc": ("cc", "CeCo", "CCOSTO", "Centro de Costo"),
    "status": ("STATUS_ EMP (ACTIVO/BAJA)", "Estatus", "Status"),
    "terminated": ("F_BAJA", "Fecha de Baja"), "joined": ("F_INGRESO", "Fecha de Ingreso"),
    "birth": ("F.NAC", "Fecha de Nacimiento"), "role": ("NOM_PUESTO", "Puesto"), "sex": ("SEXO", "Género"),
}


def query_cc(value: Any, header: str) -> str:
    text = str(value or "").strip().removesuffix(".0")
    if normalize(header) in {"CCOSTO", "CENTRO DE COSTO"} and re.fullmatch(r"\d{7,8}", text):
        # Esta exportación de RH codifica distrito (3 dígitos) + CeCo (5).
        # Sólo esta columna admite el formato compuesto; el cruce posterior
        # exige que el CeCo exista en el Directorio de tiendas abiertas.
        return text.zfill(8)[-5:]
    return clean_cc(value)


def summarize_query(sheet_name: str, raw_headers: list, rows: Iterable[tuple], valid_cc: Iterable[str], excluded_cc: Iterable[str], warnings: list[str], *, epoch=None, reference_date: date | None = None) -> tuple[dict, dict]:
    columns = resolve_columns(raw_headers, QUERY_COLUMNS, {"employee", "cc"}, "Query")
    if not {"status", "terminated"}.intersection(columns): raise ValueError("Query requiere Estatus o F_BAJA para comprobar actividad")
    rows = list(rows); valid_cc = set(valid_cc); excluded_cc = set(excluded_cc)
    cuts = set()
    for row in [tuple(raw_headers), *rows[:20]]:
        for index,value in enumerate(row[:-1]):
            if normalize(value) in {"CORTE DE INFORMACION", "FECHA DE CORTE"}:
                cut = parse_date(row[index+1], epoch)
                if cut: cuts.add(cut)
    if len(cuts) > 1: raise ValueError("Query contiene fechas de corte contradictorias")
    as_of = reference_date or (next(iter(cuts)) if cuts else datetime.now(UTC).date())
    if not cuts and reference_date is None: warnings.append("Query: sin fecha de corte; se usa la fecha de construcción como referencia.")
    optional_missing = sorted({"joined", "birth", "role", "sex"} - columns.keys())
    if optional_missing: warnings.append(f"Query: campos opcionales ausentes ({', '.join(optional_missing)}); sus resultados quedan en blanco.")
    unmatched = Counter(); excluded = Counter(); invalid_dates = Counter(); records = {}; conflicts = set(); matched = set(); duplicates = 0; cc_modes = Counter()
    def get(row, key): return row[columns[key]] if key in columns and columns[key] < len(row) else None
    def has(value): return value is not None and normalize(value) not in {"", "NA", "N A", "NULL", "NONE"}
    for row in rows:
        if not any(has(value) for value in row[:min(len(row),10)]): continue
        employee = str(get(row,"employee") or "").strip().removesuffix(".0")
        cc = query_cc(get(row,"cc"), clean_header(raw_headers[columns["cc"]]))
        if not employee or not cc: excluded["invalidKey"] += 1; continue
        if cc not in valid_cc:
            if cc in excluded_cc: excluded["storeNotOpen"] += 1
            else: unmatched[cc] += 1
            continue
        matched.add(cc)
        joined_raw=get(row,"joined"); terminated_raw=get(row,"terminated"); birth_raw=get(row,"birth")
        joined=parse_date(joined_raw,epoch); terminated=parse_date(terminated_raw,epoch); birth=parse_date(birth_raw,epoch)
        for key,raw,parsed in (("joined",joined_raw,joined),("terminated",terminated_raw,terminated),("birth",birth_raw,birth)):
            if has(raw) and parsed is None: invalid_dates[key] += 1
        status=normalize(get(row,"status"))
        if status in {"BAJA", "INACTIVO", "INACTIVA", "BAJAS"}: excluded["inactive"] += 1; continue
        if has(terminated_raw) and terminated is None: excluded["unknownActivity"] += 1; continue
        if terminated and terminated <= as_of: excluded["inactive"] += 1; continue
        if "status" in columns and status not in {"ACTIVO", "ACTIVA", "ACTIVOS", "ACTIVAS"} and (status or "terminated" not in columns):
            excluded["unknownActivity"] += 1; continue
        if joined and joined > as_of: excluded["notStartedAtCutoff"] += 1; continue
        if birth and (birth > as_of or as_of.year-birth.year > 120): invalid_dates["birth"] += 1; birth=None
        sex=normalize(get(row,"sex")); sex={"FEMENINO":"F", "MUJER":"F", "MASCULINO":"M", "HOMBRE":"M"}.get(sex,sex)
        record=(cc, clean_header(get(row,"role")), sex if sex in {"F","M"} else "", birth, joined)
        if employee in records:
            if records[employee] == record: duplicates += 1
            else: conflicts.add(employee)
        else: records[employee]=record
        raw_cc=str(get(row,"cc") or "").strip().removesuffix(".0")
        cc_modes["district3+cc5" if len(raw_cc) in {7,8} else "cc5"] += 1
    for employee in conflicts: records.pop(employee,None)
    groups=defaultdict(list)
    for record in records.values(): groups[record[0]].append(record)
    partners={}
    for cc in matched:
        active=groups[cc]; roles=Counter(record[1] for record in active if record[1]); sexes=Counter(record[2] for record in active if record[2])
        ages=[]; tenures=[]; birthdays=0; anniversaries=0
        for _,role,sex,birth,joined in active:
            if birth:
                ages.append(as_of.year-birth.year-((as_of.month,as_of.day)<(birth.month,birth.day)))
                birthdays += int(birth.month==as_of.month)
            if joined:
                tenures.append((as_of.year-joined.year)*12+as_of.month-joined.month-int(as_of.day<joined.day))
                anniversaries += int(joined.month==as_of.month)
        role_count=sum(roles.values()); gender_count=sum(sexes.values())
        partners[cc]={"headcount":len(active),
            "baristas":sum(value for key,value in roles.items() if "BARISTA" in normalize(key)) if "role" in columns else None,
            "supervisors":sum(value for key,value in roles.items() if "SUPERVISOR" in normalize(key)) if "role" in columns else None,
            "managers":sum(value for key,value in roles.items() if any(term in normalize(key) for term in ("GERENTE","STORE MANAGER"))) if "role" in columns else None,
            "female":sexes.get("F",0) if "sex" in columns else None, "male":sexes.get("M",0) if "sex" in columns else None,
            "genderKnownCount":gender_count, "roleKnownCount":role_count,
            "ageCount":len(ages), "ageTotal":sum(ages), "tenureCount":len(tenures), "tenureTotalMonths":sum(tenures),
            "avgAge":round_number(sum(ages)/len(ages),1) if ages else None,
            "avgTenureMonths":round_number(sum(tenures)/len(tenures),1) if tenures else None,
            "birthdaysThisMonth":birthdays if ages or not active and "birth" in columns else None,
            "anniversariesThisMonth":anniversaries if tenures or not active and "joined" in columns else None,
            "asOf":as_of.isoformat(), "roles":dict(roles)}
    if conflicts: warnings.append(f"Query: {len(conflicts)} empleados con asignaciones contradictorias se excluyeron de los conteos.")
    if invalid_dates: warnings.append(f"Query: {sum(invalid_dates.values())} fechas inválidas no se usaron en cálculos.")
    return partners,{"sheet":sheet_name, "rows":len(rows), "uniqueEmployees":len(records), "activeEmployees":len(records),
        "matchedStores":len(partners), "asOf":as_of.isoformat(), "asOfSource":"source" if cuts else "referenceDate",
        "activityRule":"F_BAJA + F_INGRESO al corte" if "terminated" in columns else "Estatus activo exacto",
        "headers":{key:clean_header(raw_headers[index]) for key,index in columns.items()}, "ccFormats":dict(cc_modes),
        "optionalMissingFields":optional_missing, "unmatched":dict(unmatched), "excludedRows":dict(excluded),
        "duplicateRows":duplicates, "ambiguousEmployees":len(conflicts), "invalidDates":dict(invalid_dates)}


def number(value: Any, *, percent: bool = False) -> float | None:
    if value is None or isinstance(value, bool): return None
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = str(value).strip()
        if not text or normalize(text) in {"NA", "N A", "N/A", "NULL"}: return None
        negative = text.startswith("(") and text.endswith(")")
        text = text.strip("()").replace("$", "").replace(",", "").replace("%", "")
        try: result = float(text)
        except ValueError: return None
        if negative: result = -result
        if percent: result /= 100
    return result if math.isfinite(result) else None


def round_number(value: float | None, digits: int = 6) -> float | None:
    if value is None: return None
    return round(value, digits)


def validate_xlsx(path: Path) -> None:
    if not path.is_file() or path.suffix.casefold() != ".xlsx" or not zipfile.is_zipfile(path):
        raise ValueError(f"XLSX inválido: {path.name}")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        total = sum(entry.file_size for entry in entries)
        if total > MAX_XLSX_UNCOMPRESSED: raise ValueError(f"XLSX excede límite seguro: {path.name}")
        for entry in entries:
            if entry.filename.startswith(("/", "\\")) or ".." in Path(entry.filename).parts:
                raise ValueError(f"Ruta interna insegura en {path.name}")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as output:
            json.dump(payload, output, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
            output.flush(); os.fsync(output.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def load_csv(path: Path, encoding: str = "utf-8-sig") -> list[dict[str, str]]:
    with path.open(encoding=encoding, newline="") as source:
        reader = csv.DictReader(source)
        headers = [normalize(header) for header in reader.fieldnames or []]
        if not headers or not all(headers) or len(headers) != len(set(headers)):
            raise ValueError(f"Encabezados vacíos o duplicados en {path.name}")
        rows = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"Fila CSV incompleta en {path.name}, línea {reader.line_num}")
            rows.append(row)
        return rows


def load_export_csv(path: Path) -> list[dict[str, str]]:
    raw = path.read_bytes()
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    try: text = raw.decode(encoding)
    except UnicodeError as error: raise ValueError(f"Codificación inválida en {path.name}; usa UTF-8 o UTF-16 con BOM") from error
    for delimiter in (",", ";", "\t"):
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        found_header = False
        try:
            for fields in reader:
                headers = [clean_header(value) for value in fields]
                if not {"Mes", "Tiendas"}.issubset(headers): continue
                found_header = True
                if len(headers) != len(set(headers)) or not all(headers):
                    raise ValueError(f"Encabezados vacíos o duplicados en {path.name}")
                rows = []
                for fields in reader:
                    if not any(value.strip() for value in fields): continue
                    if len(fields) != len(headers):
                        raise ValueError(f"Fila CSV incompleta en {path.name}, línea {reader.line_num}")
                    rows.append(dict(zip(headers, fields)))
                return rows
        except csv.Error as error:
            if not found_header: continue
            raise ValueError(f"CSV inválido en {path.name}: {error}") from error
    raise ValueError(f"No se encontró encabezado Mes + Tiendas en {path.name}")


def month_from_profile(value: Any) -> int | None:
    text = str(value or "").strip()
    if re.fullmatch(r"\d+(?:\.0)?", text):
        numeric = int(float(text))
        month = numeric % 100 if numeric >= 100000 else numeric
        return month if 1 <= month <= 12 else None
    match = re.match(r"^(?:\d+_)?([A-Za-zÁÉÍÓÚáéíóú]+)$", text)
    return MONTHS.get(normalize(match[1]).lower()) if match else None


def month_from_period(value: Any) -> int | None:
    text = str(value or "").strip()
    if not re.fullmatch(r"20\d{4}", text): return None
    month = int(text[-2:])
    return month if 1 <= month <= 12 else None


def year_from_period(value: Any) -> int | None:
    text = str(value or "").strip()
    return int(text[:4]) if month_from_period(text) is not None else None


def load_business(paths: dict[str, Path], valid_cc: Iterable[str], warnings: list[str]) -> tuple[dict, dict]:
    valid_cc = set(valid_cc)
    inputs = {name: load_export_csv(paths[name]) for name in ("business_aa", "business_real")}
    required = {
        "business_aa": {"Mes", "Tiendas", "ADT AA", "OMT"},
        "business_real": {"Mes", "Tiendas", "ADT Real", "Venta $", "Var Ventas vs Ppto (%)", "AWS $", "Ticket Prom Real", "Ticket Prom AA", "Var Ticket vs AA (%)"},
    }
    for name, rows in inputs.items():
        headers = set(rows[0]) if rows else set()
        if missing := required[name] - headers:
            raise ValueError(f"{paths[name].name} sin encabezados requeridos: {sorted(missing)}")
    real_headers = list(inputs["business_real"][0])
    optional_missing = sorted({"Ticket Prom Ppto", "Var Ticket vs Ppto (%)"} - set(real_headers))
    if optional_missing:
        warnings.append(f"Ticket: referencias opcionales ausentes ({', '.join(optional_missing)}); se muestran en blanco.")

    years = Counter(year_from_period(row.get("Mes")) for rows in inputs.values() for row in rows
                    if year_from_period(row.get("Mes")) is not None and clean_cc(row.get("Tiendas")) in valid_cc)
    if not years: raise ValueError("Negocio no contiene periodos válidos para tiendas del Directorio")
    selected_year = max(years)
    ignored = Counter(); unmatched = Counter(); months = Counter(); seen = set()
    business = defaultdict(dict)
    for name, rows in inputs.items():
        for row in rows:
            month = month_from_period(row.get("Mes")); year = year_from_period(row.get("Mes")); cc = clean_cc(row.get("Tiendas"))
            if month is None or not cc: continue
            if cc not in valid_cc: unmatched[cc] += 1; continue
            if year != selected_year: ignored[year] += 1; continue
            key = (cc, month, name)
            if key in seen: raise ValueError(f"Negocio duplicado: {cc}, periodo {year}{month:02d}, motor {name}")
            seen.add(key)
            target = business[cc].setdefault(str(month), {})
            if name == "business_aa":
                target.update({"adtAa": round_number(number(row.get("ADT AA"))), "omtDiff": round_number(number(row.get("OMT")))})
            else:
                sales = number(row.get("Venta $")); variance = number(row.get("Var Ventas vs Ppto (%)"), percent=True)
                budget = sales / (1 + variance) if sales is not None and variance is not None and not math.isclose(variance, -1) else None
                target.update({
                    "adt": round_number(number(row.get("ADT Real"))), "sales": round_number(sales, 2),
                    "salesBudget": round_number(budget, 2), "salesVariance": round_number(variance),
                    "aws": round_number(number(row.get("AWS $")), 2), "ticket": round_number(number(row.get("Ticket Prom Real")), 2),
                    "ticketAa": round_number(number(row.get("Ticket Prom AA")), 2),
                    "ticketBudget": round_number(number(row.get("Ticket Prom Ppto")), 2),
                    "ticketVariance": round_number(number(row.get("Var Ticket vs AA (%)"), percent=True)),
                    "ticketBudgetVariance": round_number(number(row.get("Var Ticket vs Ppto (%)"), percent=True)),
                })
            months[month] += 1
    if ignored: warnings.append(f"Negocio: se seleccionó {selected_year}; {sum(ignored.values())} filas de otros años no se mezclaron.")
    return business, {"matchedStores": len(business), "months": dict(months), "years": dict(years),
                      "selectedYear": selected_year, "ignoredYears": dict(ignored), "unmatched": dict(unmatched),
                      "realHeaders": real_headers, "optionalMissingHeaders": optional_missing}


def metric_format(header: str) -> str:
    key = normalize(header)
    if key in {"IPLH", "IPLH AA", "TPLH", "TPLH AA"}: return "decimal"
    if key in {"ICA SCORE", "ICA SCORE AA", "SEGUNDAS CONEXIONES", "SEGUNDAS CONEXIONES AA"}: return "number"
    if key in {"DT TIME", "TIEMPO DT AA"}: return "duration"
    return "percent"


def metric_direction(graph: str) -> str:
    key = normalize(graph)
    return "lower" if key in {"LABOR", "COSTO", "COSTO %", "DT TIME", "ROTACION"} else "higher"


def pair_kind(reference: str) -> str:
    return "ppto" if "PPTO" in normalize(reference) else "aa"


def build(root: Path, output: Path, audit_output: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    engines = root / "data" / "engines"
    paths = {
        "directory": engines / "Directorio_Perfil Tienda.csv",
        "profile": engines / "Base_Perfil Tienda.xlsx",
        "business_aa": engines / "Base_Perfil Tienda.csv",
        "business_real": engines / "Base_Perfil Tienda_2.csv",
        "mix_manifest": engines / "mix" / "manifest.json",
        "partners": engines / "Query.xlsx",
    }
    for path in paths.values():
        if not path.is_file(): raise ValueError(f"Falta motor: {path.name}")
    validate_xlsx(paths["profile"]); validate_xlsx(paths["partners"])

    issues: list[str] = []
    warnings: list[str] = []
    sources = {name: {"file": str(path.relative_to(engines)), "bytes": path.stat().st_size, "sha256": sha256(path)} for name, path in paths.items()}
    with paths["mix_manifest"].open(encoding="utf-8") as source:
        mix_manifest = json.load(source)
    mix_paths = [paths["mix_manifest"].parent / part["file"] for part in mix_manifest.get("parts", [])]
    if not mix_paths or any(not path.is_file() for path in mix_paths):
        raise ValueError("El manifiesto de Base_Mix no contiene todas sus partes")
    for path, part in zip(mix_paths, mix_manifest["parts"]):
        if path.stat().st_size != part["bytes"] or sha256(path) != part["sha256"]:
            raise ValueError(f"Parte Mix alterada o incompleta: {path.name}")
        sources[f"mix_{part['month']:02d}"] = {"file": str(path.relative_to(engines)), "bytes": path.stat().st_size, "sha256": part["sha256"]}

    directory, aliases, directory_audit = load_directory(paths["directory"])
    directory_by_cc = {item["cc"]: item for item in directory}
    excluded_cc = set(directory_audit["excludedCecos"])

    workbook = load_workbook(paths["profile"], read_only=True, data_only=True, keep_links=False)
    if "Perfil" not in workbook.sheetnames or "Instrucciones_Ejemplo" not in workbook.sheetnames:
        raise ValueError("Base_Perfil Tienda.xlsx requiere Perfil e Instrucciones_Ejemplo")
    profile_sheet = workbook["Perfil"]
    headers = [clean_header(value) for value in next(profile_sheet.iter_rows(values_only=True))]
    normalized_headers = [normalize(header) for header in headers]
    month_aliases = {"MES", "MES NUM", "MES NUMERO", "PERIODO"}
    try: month_column = next(index for index, header in enumerate(normalized_headers) if header in month_aliases)
    except StopIteration as error: raise ValueError("Perfil requiere MES_NUM, Mes o Periodo") from error
    try: cc_column = next(index for index, header in enumerate(normalized_headers) if header in {"CECO", "CC", "TIENDA", "TIENDAS"})
    except StopIteration as error: raise ValueError("Perfil requiere CeCo") from error
    if month_column == cc_column: raise ValueError("Mes y CeCo no pueden compartir columna")
    duplicate_headers = [header for header, count in Counter(headers).items() if count > 1]
    if duplicate_headers: raise ValueError(f"Encabezados duplicados en Perfil: {duplicate_headers}")

    instruction_rows = []
    for row in workbook["Instrucciones_Ejemplo"].iter_rows(min_row=2, values_only=True):
        values = [clean_header(value) for value in row[:5]]
        if any(values): instruction_rows.append({"pillar": values[0], "header": values[1], "graph": values[2], "instruction": values[3], "note": values[4]})
    instruction_by_header = {normalize(item["header"]): item for item in instruction_rows if item["header"]}

    metric_columns = [index for index in range(len(headers)) if index not in {month_column, cc_column}]
    metric_headers = [headers[index] for index in metric_columns]
    metric_index = {header: index for index, header in enumerate(metric_headers)}
    graphs: list[dict[str, Any]] = []
    seen_graphs: set[tuple[str, str]] = set()
    visible_instructions = [item for item in instruction_rows if normalize(item["graph"]) not in {"", "NO SE MUESTRA"} and normalize(item["pillar"]) in {"PARTNER", "CLIENTE", "NEGOCIO"}]
    header_aliases = {"ROTACION": "ROLLING RY", "ROTACION AA": "ROLLING RY AA"}
    for item in visible_instructions:
        graph_key = (normalize(item["pillar"]), normalize(item["graph"]))
        if graph_key in seen_graphs: continue
        seen_graphs.add(graph_key)
        candidates = [entry["header"] for entry in visible_instructions if (normalize(entry["pillar"]), normalize(entry["graph"])) == graph_key]
        actual = next((header for header in candidates if " AA" not in f" {normalize(header)}" and "PPTO" not in normalize(header)), candidates[0])
        reference = next((header for header in candidates if normalize(header) != normalize(actual)), "")
        actual_key = header_aliases.get(normalize(actual), normalize(actual))
        reference_key = header_aliases.get(normalize(reference), normalize(reference))
        actual = next((header for header in metric_headers if normalize(header) == actual_key), actual)
        reference = next((header for header in metric_headers if normalize(header) == reference_key), reference)
        if actual not in metric_index:
            warnings.append(f"Métrica clasificada no encontrada: {actual}")
            continue
        graph_name = normalize(item["graph"])
        metric_name = normalize(actual)
        # VMT y OMT son indicadores operativos de Negocio aunque la hoja de
        # clasificación de versiones anteriores los ubicara en Cliente.
        pillar = "Negocio" if graph_name in {"VMT", "OMT"} or metric_name in {"VMT", "OMT"} else item["pillar"]
        graphs.append({
            "id": re.sub(r"[^a-z0-9]+", "-", normalize(item["graph"]).lower()).strip("-"),
            "pillar": pillar, "title": item["graph"], "actual": actual,
            "reference": reference if reference in metric_index else None,
            "referenceKind": pair_kind(reference), "format": metric_format(actual),
            "direction": metric_direction(item["graph"]),
            "ytd": "latest" if normalize(item["graph"]) == "ROTACION" else "average",
            "instruction": item["instruction"], "note": item["note"],
        })

    profile: defaultdict[str, dict[str, list[float | None]]] = defaultdict(dict)
    profile_seen: set[tuple[str, int]] = set()
    profile_unmatched = Counter(); profile_excluded = Counter(); minus_100_blanked = Counter(); duration_blanked = Counter(); profile_months = Counter()
    for row_number, row in enumerate(profile_sheet.iter_rows(min_row=2, values_only=True), start=2):
        month = month_from_profile(row[month_column]); cc = clean_cc(row[cc_column])
        if month is None or not cc:
            warnings.append(f"Perfil fila {row_number}: periodo o CeCo inválido")
            continue
        if cc not in directory_by_cc:
            if cc in excluded_cc: profile_excluded[cc] += 1; continue
            profile_unmatched[cc] += 1; continue
        key = (cc, month)
        if key in profile_seen:
            issues.append(f"Perfil duplicado: {cc}, mes {month}"); continue
        profile_seen.add(key); profile_months[month] += 1
        values: list[float | None] = []
        for header, column in zip(metric_headers, metric_columns):
            raw = row[column] if column < len(row) else None
            value = number(raw)
            if metric_format(header) == "percent" and value is not None and math.isclose(value, -1.0, abs_tol=1e-9):
                minus_100_blanked[header] += 1; value = None
            if metric_format(header) == "duration" and value is not None and abs(value) < 1:
                # El motor exporta DT como fracción de día que Excel interpreta
                # visualmente como hh:mm. El negocio lo define como mm:ss:
                # 0.052083 (01:15 en origen) debe representar 75 segundos.
                value *= 1440
            if metric_format(header) == "duration" and value is not None and not 20 <= value <= 1800:
                duration_blanked[header] += 1; value = None
            values.append(round_number(value))
        profile[cc][str(month)] = values

    business, business_audit = load_business(paths, directory_by_cc, warnings)
    business_audit["excludedClosedRows"] = sum(business_audit["unmatched"].pop(cc,0) for cc in excluded_cc)

    alias_to_cc = {alias: next(iter(ccs)) for alias, ccs in aliases.items() if len(ccs) == 1}
    mix: defaultdict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    mix_unmatched = Counter(); mix_excluded = 0; mix_invalid_sales = 0; mix_rows = 0; mix_matched_rows = 0
    mix_months = Counter()
    for mix_path in mix_paths:
        with mix_path.open(encoding="utf-8-sig", newline="") as source:
            for row in csv.DictReader(source):
                mix_rows += 1; month = month_from_profile(row.get("Mes"))
                cc = alias_to_cc.get(normalize(row.get("Tienda")))
                if cc in excluded_cc: mix_excluded += 1; continue
                if not cc:
                    mix_unmatched[clean_header(row.get("Tienda"))] += 1; continue
                sale = number(row.get("Venta"))
                if month is None or sale is None:
                    mix_invalid_sales += 1; continue
                mix_matched_rows += 1; mix_months[month] += 1
                target = mix[cc].setdefault(str(month), {"category": defaultdict(float), "order": defaultdict(float), "total": 0.0})
                category = clean_header(row.get("Category")) or "Sin categoría"
                order = clean_header(row.get("Tipo Orden")) or "Sin canal"
                target["category"][category] += sale; target["order"][order] += sale; target["total"] += sale
    for months in mix.values():
        for target in months.values():
            total = target["total"]
            target["category"] = {key: round_number(value / total) for key, value in target["category"].items()} if total else {}
            target["order"] = {key: round_number(value / total) for key, value in target["order"].items()} if total else {}
            target["total"] = round_number(total, 4)

    partner_book = load_workbook(paths["partners"], read_only=True, data_only=True, keep_links=False)
    try:
        query_names = [name for name in partner_book.sheetnames if normalize(name) == "QUERY"]
        if len(query_names) != 1: raise ValueError("Query.xlsx requiere una única pestaña Query (sin distinguir mayúsculas)")
        instructions_name = next((name for name in partner_book.sheetnames if normalize(name) == "INSTRUCCIONES"), None)
        if instructions_name:
            for row in partner_book[instructions_name].iter_rows(values_only=True):
                values = [clean_header(value) for value in row[:2]]
                if any(values): instruction_rows.append({"pillar":"Query Partner", "header":values[0], "graph":"Uso", "instruction":values[1] if len(values)>1 else "", "note":""})
        query = partner_book[query_names[0]]
        query_headers = list(next(query.iter_rows(values_only=True)))
        partners, partner_audit = summarize_query(query.title, query_headers, query.iter_rows(min_row=2, values_only=True),
            directory_by_cc, excluded_cc, warnings, epoch=partner_book.epoch)
    finally:
        partner_book.close()

    all_months = sorted({int(month) for values in profile.values() for month in values} | {int(month) for values in business.values() for month in values} | {int(month) for values in mix.values() for month in values})
    coverage = []
    for item in directory:
        cc = item["cc"]
        coverage.append({"cc": cc, "profile": cc in profile, "business": cc in business, "mix": cc in mix, "partners": cc in partners})
    for count, message in (
        (len(profile_unmatched), "Perfil: {count} CeCo sin cruce; se dejaron en blanco."),
        (len(business_audit["unmatched"]), "Negocio: {count} CeCo sin cruce; se dejaron en blanco."),
        (len(mix_unmatched), "Mix: {count} nombres sin coincidencia exacta única; se dejaron en blanco."),
        (len(partner_audit["unmatched"]), "Query: {count} CeCo sin cruce; se dejaron en blanco."),
    ):
        if count: warnings.append(message.format(count=count))
    if directory_audit["invalidOpeningDates"]: warnings.append(f"Directorio: {directory_audit['invalidOpeningDates']} fechas de apertura inválidas quedaron en blanco.")

    audit = {
        "schemaVersion": 2, "generatedAt": datetime.now(UTC).isoformat(), "issueCount": len(issues), "warningCount": len(warnings),
        "issues": issues, "warnings": warnings, "sources": sources,
        "directory": directory_audit,
        "profile": {"rows": profile_sheet.max_row - 1, "matchedStores": len(profile), "months": dict(profile_months), "monthHeader": headers[month_column], "ccHeader": headers[cc_column], "minus100Blanked": dict(minus_100_blanked), "durationOutliersBlanked": dict(duration_blanked), "unmatched": dict(profile_unmatched), "excludedClosedRows":sum(profile_excluded.values())},
        "business": business_audit,
        "mix": {"sourceRows": mix_rows, "manifestRows": mix_manifest.get("rows"), "parts": len(mix_paths), "months": dict(mix_months), "matchedRows": mix_matched_rows, "matchedStores": len(mix), "invalidRows": mix_invalid_sales, "excludedClosedRows":mix_excluded, "unmatchedNames": len(mix_unmatched), "unmatchedTop": dict(mix_unmatched.most_common(100))},
        "partners": partner_audit,
        "coverage": coverage,
    }
    if issues: raise ValueError(" | ".join(issues))

    display_year = business_audit["selectedYear"]
    payload = {
        "schemaVersion": 2, "generatedAt": audit["generatedAt"], "directoryPolicy":"open-only-v1",
        "months": [{"id": month, "period": f"{display_year}{month:02d}", "label": MONTH_LABELS[month - 1], "short": MONTH_LABELS[month - 1][:3]} for month in all_months],
        "directory": directory, "metricHeaders": metric_headers, "graphs": graphs,
        "profile": profile, "business": business, "mix": mix, "partners": partners,
        "engineInfo": {"profile":{"months":sorted(profile_months)}, "business":{"months":sorted(business_audit["months"]),"year":display_year},
                       "mix":{"months":sorted(mix_months)}, "partners":{"asOf":partner_audit["asOf"],"asOfSource":partner_audit["asOfSource"]}},
        "instructions": instruction_rows, "auditSummary": {"warnings": len(warnings), "generatedAt": audit["generatedAt"]},
    }
    atomic_json(output, payload); atomic_json(audit_output, audit)
    return payload, audit


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path)
    parser.add_argument("--audit-output", type=Path)
    args = parser.parse_args(); root = args.root.resolve()
    output = args.output.resolve() if args.output else root / "data" / "dashboard.json"
    audit_output = args.audit_output.resolve() if args.audit_output else root / "data" / "audit.json"
    try: payload, audit = build(root, output, audit_output)
    except (OSError, ValueError, zipfile.BadZipFile) as error: raise SystemExit(f"Construcción cancelada: {error}") from error
    print(json.dumps({"status": "ready", "stores": len(payload["directory"]), "months": len(payload["months"]), "warnings": audit["warningCount"], "output": str(output)}, ensure_ascii=False))


if __name__ == "__main__": main()
