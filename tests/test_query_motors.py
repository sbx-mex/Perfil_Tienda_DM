from __future__ import annotations

import csv
import json
import sys
import tempfile
import unittest
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_data import as_iso, load_directory, query_cc, resolve_columns, summarize_query

CUT = date(2026, 8, 30)


class QueryMotorTests(unittest.TestCase):
    def directory(self, rows):
        temporary = tempfile.TemporaryDirectory(); self.addCleanup(temporary.cleanup)
        path = Path(temporary.name) / "directory.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as source:
            writer = csv.DictWriter(source, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
        return load_directory(path)

    def query(self, headers, rows, stores=("38101",), excluded=()):
        warnings = []
        data, audit = summarize_query("query", headers, rows, stores, excluded, warnings, reference_date=CUT)
        return data, audit, warnings

    def test_directory_only_accepts_open_status_and_new_store_header(self):
        rows = [{"CC":"38101", "CC Nombre":"Ángel", "Estatus":" Abierta ", "División ":"Centro", "Fecha de Apertura":"37505"}]
        for index,status in enumerate(("Cierre Temporal", "Cierre Definitivo", "Próxima Apertura", "", "No Abierta"),start=2):
            rows.append({**rows[0], "CC":str(38100+index), "Estatus":status})
        data, aliases, audit = self.directory(rows)
        self.assertEqual([item["cc"] for item in data], ["38101"])
        self.assertEqual(data[0]["store"], "Ángel")
        self.assertEqual(data[0]["division"], "Centro")
        self.assertEqual(data[0]["opened"], "2002-09-06")
        self.assertEqual(audit["excludedStores"],5)

    def test_directory_cannot_guess_status_or_duplicate_open_store(self):
        with self.assertRaisesRegex(ValueError,"status"):
            self.directory([{"CC":"38101", "Tienda":"Prueba"}])
        row = {"CC":"38101", "Tienda":"Prueba", "Estatus":"Abierta"}
        with self.assertRaisesRegex(ValueError,"duplicado"): self.directory([row,row])

    def test_date_reader_handles_excel_and_real_dates_without_inventing_dates(self):
        for raw in (37505, "37505", "06/09/2002", "2002-09-06", datetime(2002,9,6)):
            self.assertEqual(as_iso(raw),"2002-09-06")
        for raw in ("31/02/2026", "sin fecha", None): self.assertIsNone(as_iso(raw))
        self.assertEqual(as_iso("27/09/26"),"2026-09-27")
        self.assertEqual(as_iso("25 /09/2026"),"2026-09-25")

    def test_composite_cc_is_only_allowed_in_the_hr_column(self):
        self.assertEqual(query_cc("01738101", "CCOSTO"), "38101")
        self.assertEqual(query_cc(1938262, "CCOSTO"), "38262")
        self.assertEqual(query_cc("38101", "cc"), "38101")
        self.assertEqual(query_cc("01738101", "cc"), "")
        self.assertEqual(query_cc("99901738101", "CCOSTO"), "")

    def test_reduced_query_keeps_headcount_without_fabricating_demographics(self):
        data,audit,warnings = self.query(["NUM_EMP", "CCOSTO", "F_BAJA"], [(101,"01738101",None)])
        row=data["38101"]
        self.assertEqual(row["headcount"],1)
        for key in ("baristas", "female", "avgAge", "avgTenureMonths", "birthdaysThisMonth", "anniversariesThisMonth"):
            self.assertIsNone(row[key],key)
        self.assertEqual(set(audit["optionalMissingFields"]), {"role","sex","birth","joined"})
        self.assertTrue(warnings)

    def test_legacy_inactive_status_cannot_be_mistaken_for_active(self):
        data,audit,_ = self.query(["NUM_EMP", "cc", "STATUS_ EMP (ACTIVO/BAJA)"], [(101,"38101","INACTIVO"),(102,"38101","ACTIVO"),(103,"38101","desconocido")])
        self.assertEqual(data["38101"]["headcount"],1)
        self.assertEqual(audit["excludedRows"],{"inactive":1,"unknownActivity":1})

    def test_source_cutoff_controls_termination_ingress_and_age(self):
        headers=["NUM_EMP","CCOSTO","F_BAJA","F_INGRESO","F.NAC","CORTE DE INFORMACIÓN",datetime(2026,8,30)]
        rows=[(101,"01738101",None,datetime(2026,8,1),datetime(2000,8,31)),
              (102,"01738101",datetime(2026,8,30),datetime(2020,1,1),None),
              (103,"01738101",None,datetime(2026,8,31),None),
              (104,"01738101",datetime(2026,9,1),datetime(2020,1,1),None),
              (105,"01738101","fecha inválida",None,None)]
        data,audit=summarize_query("query",headers,rows,{"38101"},set(),[],reference_date=None)
        self.assertEqual(audit["asOf"],"2026-08-30")
        self.assertEqual(data["38101"]["headcount"],2)
        self.assertEqual(data["38101"]["avgAge"],25)
        self.assertEqual(data["38101"]["ageCount"],1)
        self.assertEqual(audit["excludedRows"],{"inactive":1,"notStartedAtCutoff":1,"unknownActivity":1})

    def test_closed_or_unverified_store_never_receives_partner_data(self):
        data,audit,_=self.query(["NUM_EMP","CCOSTO","F_BAJA"],[(101,"01738102",None),(102,"01799999",None)],excluded={"38102"})
        self.assertEqual(data,{})
        self.assertEqual(audit["excludedRows"],{"storeNotOpen":1})
        self.assertEqual(audit["unmatched"],{"99999":1})

    def test_repetitions_are_deduplicated_and_conflicting_assignments_are_excluded(self):
        data,audit,warnings=self.query(["NUM_EMP","CCOSTO","F_BAJA"],[(101,"01738101",None),(101,"01738101",None),(102,"01738101",None),(102,"01738103",None)],stores={"38101","38103"})
        self.assertEqual(sum(row["headcount"]for row in data.values()),1)
        self.assertEqual(audit["duplicateRows"],1)
        self.assertEqual(audit["ambiguousEmployees"],1)
        self.assertTrue(any("contradictorias"in warning for warning in warnings))

    def test_columns_with_ambiguous_meaning_are_rejected(self):
        with self.assertRaisesRegex(ValueError,"ambiguas"):
            resolve_columns(["cc", "CCOSTO"], {"cc":("cc","CCOSTO")}, {"cc"}, "Query")
        with self.assertRaisesRegex(ValueError,"Estatus o F_BAJA"):
            self.query(["NUM_EMP","CCOSTO"],[(101,"01738101")])

    def test_raw_employee_details_are_not_emitted(self):
        data,audit,_=self.query(["NUM_EMP","NOMBRE","CCOSTO","F_BAJA","F.NAC"],[(123456789,"Nombre privado de prueba","01738101",None,datetime(2000,1,1))])
        public=json.dumps({"data":data,"audit":audit})
        for value in ("Nombre privado de prueba", "123456789", "2000-01-01"):
            self.assertNotIn(value,public)

    def test_current_sources_reconcile_to_the_open_store_and_partner_population(self):
        directory,_,audit=load_directory(ROOT/"data/engines/Directorio_Perfil Tienda.csv")
        stores={row["cc"] for row in directory}
        with (ROOT/"data/engines/Directorio_Perfil Tienda.csv").open(encoding="utf-8-sig",newline="")as source:
            expected={row["CC"].zfill(5) for row in csv.DictReader(source) if row["Estatus"].strip().casefold()=="abierta"}
        self.assertEqual(stores,expected)
        book=load_workbook(ROOT/"data/engines/Query.xlsx",read_only=True,data_only=True)
        try:
            sheet=book[next(name for name in book.sheetnames if name.casefold()=="query")]
            headers=list(next(sheet.iter_rows(values_only=True))); rows=list(sheet.iter_rows(min_row=2,values_only=True))
            data,meta=summarize_query(sheet.title,headers,rows,stores,audit["excludedCecos"],[],epoch=book.epoch)
            index={str(header):i for i,header in enumerate(headers)}
            cut=headers[headers.index("CORTE DE INFORMACIÓN")+1].date()
            counts=Counter()
            for row in rows:
                cc=str(row[index["CCOSTO"]]).removesuffix(".0")[-5:]
                if cc in expected and row[index["F_BAJA"]] is None and row[index["F_INGRESO"]].date()<=cut:
                    counts[cc]+=1
            self.assertEqual(sum(row["headcount"]for row in data.values()),sum(counts.values()))
            self.assertEqual({cc:row["headcount"]for cc,row in data.items() if row["headcount"]},dict(counts))
            self.assertEqual(meta["asOf"],cut.isoformat())
        finally: book.close()


if __name__=="__main__":unittest.main()
